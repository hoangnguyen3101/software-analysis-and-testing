"""Tao tai lieu on tap Lab03 dang Word (.docx).

Doc cac hinh trong Lab03/docs/images (sinh boi tools/figures.py) va ghi ra
Lab03/docs/Lab03-Tai-lieu-on-tap.docx.

Chay (tu thu muc Lab03):
    pip install -r tools/requirements.txt
    python tools/figures.py      # ve lai hinh truoc
    python tools/build_docx.py   # hoac: python tools/build_docx.py <duong-dan-file.docx>
"""

import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

DOCS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "docs")
FIG = os.path.join(DOCS, "images")
OUT = sys.argv[1] if len(sys.argv) > 1 else os.path.join(DOCS, "Lab03-Tai-lieu-on-tap.docx")

BODY_FONT = "Calibri"
CODE_FONT = "Consolas"
C_H1 = RGBColor(0x1F, 0x4E, 0x79)
C_H2 = RGBColor(0x2E, 0x75, 0xB6)
C_CODE = RGBColor(0xAD, 0x14, 0x57)
HEADER_FILL = "1F4E79"
ZEBRA_FILL = "F2F6FC"

doc = Document()


# ---------------------------------------------------------------- styles

def set_run_font(run, name):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)


def style_font(style, name, size=None, bold=None, color=None):
    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rfonts.get(qn(attr)) is not None:
            del rfonts.attrib[qn(attr)]
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if size:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color


normal = doc.styles["Normal"]
style_font(normal, BODY_FONT, 12)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.15

for name, size, color, before, after in [("Heading 1", 17, C_H1, 18, 8),
                                         ("Heading 2", 14, C_H2, 14, 6),
                                         ("Heading 3", 12.5, C_H2, 10, 4)]:
    st = doc.styles[name]
    style_font(st, BODY_FONT, size, True, color)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for name in ("List Bullet", "List Bullet 2"):
    style_font(doc.styles[name], BODY_FONT, 12)
    doc.styles[name].paragraph_format.space_after = Pt(3)

sec = doc.sections[0]
sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)
sec.left_margin = Cm(2.3)
sec.right_margin = Cm(2.0)
CONTENT_W = 21.0 - 2.3 - 2.0


# ---------------------------------------------------------------- helpers

TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def add_rich(par, text, size=None, bold=False, italic=False, color=None):
    """Them doan van co dinh dang: **dam**, `ma nguon`."""
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            set_run_font(run, CODE_FONT)
            run.font.color.rgb = C_CODE
            shade(run._element.get_or_add_rPr(), "F3F3F3")
            if size:
                run.font.size = Pt(size - 1)
            else:
                run.font.size = Pt(11)
            if bold:
                run.bold = True
            continue
        else:
            run = par.add_run(part)
            if bold:
                run.bold = True
        if italic:
            run.italic = True
        if size:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = color
    return par


def H1(text):
    doc.add_heading(text, level=1)


def H2(text):
    doc.add_heading(text, level=2)


def H3(text):
    doc.add_heading(text, level=3)


def P(text, align=None, size=None, bold=False, italic=False, color=None, after=None, before=None):
    par = doc.add_paragraph()
    add_rich(par, text, size=size, bold=bold, italic=italic, color=color)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        par.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    if before is not None:
        par.paragraph_format.space_before = Pt(before)
    # doan dan dat ket thuc bang ":" phai di cung noi dung ngay sau no
    if text.rstrip().endswith(":"):
        par.paragraph_format.keep_with_next = True
    return par


def B(text, level=0):
    par = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    add_rich(par, text)
    return par


def PB():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def set_cell_fill(cell, fill):
    shade(cell._tc.get_or_add_tcPr(), fill)


def set_cell_borders(cell, **kw):
    tcpr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        spec = kw.get(edge)
        el = OxmlElement(f"w:{edge}")
        if spec is None:
            el.set(qn("w:val"), "nil")
        else:
            val, sz, color = spec
            el.set(qn("w:val"), val)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), color)
        borders.append(el)
    tcpr.append(borders)


def set_cell_margins(cell, top=80, bottom=80, left=140, right=140):
    tcpr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for k, v in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{k}")
        el.set(qn("w:w"), str(v))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcpr.append(mar)


def row_no_split(row, header=False):
    trpr = row._tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trpr.append(cs)
    if header:
        th = OxmlElement("w:tblHeader")
        trpr.append(th)


def fill_cell(cell, text, size=10.5, bold=False, color=None, align=None, font=None):
    cell.text = ""
    lines = text.split("\n")
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(1)
        par.paragraph_format.line_spacing = 1.05
        if font == "code":
            run = par.add_run(line)
            set_run_font(run, CODE_FONT)
            run.font.size = Pt(size - 0.5)
            if bold:
                run.bold = True
        else:
            add_rich(par, line, size=size, bold=bold, color=color)
        if align == "center":
            par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def T(headers, rows, widths=None, center_cols=(), size=10.5, caption=None, code_cols=(), zebra=True,
      highlight_rows=()):
    ncol = len(headers)
    table = doc.add_table(rows=1, cols=ncol)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths is None:
        widths = [CONTENT_W / ncol] * ncol
    hdr = table.rows[0]
    row_no_split(hdr, header=True)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        fill_cell(c, h, size=size, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align="center")
        set_cell_fill(c, HEADER_FILL)
    for r_i, r in enumerate(rows):
        row = table.add_row()
        row_no_split(row)
        for i, val in enumerate(r):
            c = row.cells[i]
            fill_cell(c, val, size=size, align="center" if i in center_cols else None,
                      font="code" if i in code_cols else None)
            if r_i in highlight_rows:
                set_cell_fill(c, "FFF2CC")
            elif zebra and r_i % 2 == 1:
                set_cell_fill(c, ZEBRA_FILL)
    for row in table.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    # bang ngan: giu tron tren mot trang
    if len(rows) <= 10:
        for row in table.rows[:-1]:
            for c in row.cells:
                for par in c.paragraphs:
                    par.paragraph_format.keep_with_next = True
    if caption:
        P(caption, align="center", size=10, italic=True, color=RGBColor(0x59, 0x59, 0x59), before=3)
    else:
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(2)
    return table


def CODE(text, size=9.5, title=None):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Cm(CONTENT_W)
    set_cell_fill(cell, "F6F8FA")
    grey = ("single", 6, "D0D7DE")
    set_cell_borders(cell, top=grey, bottom=grey, right=grey, left=("single", 24, "2E75B6"))
    set_cell_margins(cell, 100, 100, 200, 140)
    cell.text = ""
    lines = text.strip("\n").split("\n")
    if title:
        par = cell.paragraphs[0]
        run = par.add_run(title)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
        par.paragraph_format.space_after = Pt(3)
    for i, line in enumerate(lines):
        par = cell.paragraphs[0] if (i == 0 and not title) else cell.add_paragraph()
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.line_spacing = 1.0
        run = par.add_run(line if line else " ")
        set_run_font(run, CODE_FONT)
        run.font.size = Pt(size)
    # khoi code ngan thi giu tren mot trang, khoi dai cho phep tach trang
    if len(lines) <= 12:
        row_no_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


BOX_KINDS = {
    "note": ("DEEBF7", "2E75B6", "GHI NHỚ"),
    "warn": ("FDECEA", "C0392B", "LỖI THƯỜNG GẶP"),
    "tip": ("E8F5E9", "2E7D32", "MẸO"),
    "ex": ("FFF8E1", "F9A825", "VÍ DỤ"),
    "q": ("F3E5F5", "7B1FA2", "CÂU HỎI"),
}


def BOX(kind, items, title=None):
    fill, border, default_title = BOX_KINDS[kind]
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Cm(CONTENT_W)
    set_cell_fill(cell, fill)
    set_cell_borders(cell, left=("single", 36, border), top=None, bottom=None, right=None)
    set_cell_margins(cell, 110, 110, 220, 160)
    cell.text = ""
    par = cell.paragraphs[0]
    run = par.add_run(title or default_title)
    run.bold = True
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor.from_string(border)
    par.paragraph_format.space_after = Pt(3)
    for it in items:
        par = cell.add_paragraph()
        par.paragraph_format.space_after = Pt(2)
        if it.startswith("• "):
            par.paragraph_format.left_indent = Cm(0.5)
            par.paragraph_format.first_line_indent = Cm(-0.4)
            add_rich(par, "•  " + it[2:], size=11)
        else:
            add_rich(par, it, size=11)
    row_no_split(table.rows[0])
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def FORMULA(text, note=None, width=11, size=14):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.rows[0].cells[0]
    cell.width = Cm(width)
    set_cell_fill(cell, "FFFDE7")
    b = ("single", 8, "F9A825")
    set_cell_borders(cell, top=b, bottom=b, left=b, right=b)
    set_cell_margins(cell, 100, 100, 160, 160)
    cell.text = ""
    par = cell.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.space_after = Pt(0)
    run = par.add_run(text)
    set_run_font(run, "Cambria Math")
    run.font.size = Pt(size)
    run.bold = True
    if note:
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        add_rich(p2, note, size=10.5, italic=True)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def IMG(name, width_cm, caption):
    par = doc.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.paragraph_format.space_after = Pt(2)
    par.add_run().add_picture(os.path.join(FIG, name), width=Cm(width_cm))
    P(caption, align="center", size=10, italic=True, color=RGBColor(0x59, 0x59, 0x59), after=10)


def add_field(run, instr):
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    t = OxmlElement("w:instrText")
    t.set(qn("xml:space"), "preserve")
    t.text = instr
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    run._r.append(b)
    run._r.append(t)
    run._r.append(e)


# ---------------------------------------------------------------- header / footer

sec.different_first_page_header_footer = True
hp = sec.header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run("Lab03 — Kiểm thử hộp trắng  |  Tài liệu ôn tập")
r.font.size = Pt(9)
r.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)

fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for piece in ("Trang ", "PAGE", " / ", "NUMPAGES"):
    run = fp.add_run()
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x7F, 0x7F, 0x7F)
    if piece in ("PAGE", "NUMPAGES"):
        add_field(run, piece)
    else:
        run.text = piece

# ======================================================================
# TRANG BIA
# ======================================================================

for _ in range(5):
    doc.add_paragraph()
P("PHÂN TÍCH VÀ KIỂM THỬ PHẦN MỀM", align="center", size=13, bold=True, color=RGBColor(0x7F, 0x7F, 0x7F))
P("LAB03", align="center", size=40, bold=True, color=C_H1, after=0)
P("KIỂM THỬ HỘP TRẮNG", align="center", size=22, bold=True, color=C_H2, after=14)
P("Control Flow Graph  •  Cyclomatic Complexity  •  Đường đi độc lập\nStatement Coverage  •  Branch Coverage",
  align="center", size=13, color=RGBColor(0x40, 0x40, 0x40), after=30)
P("Tài liệu ôn tập — giải thích từng bước qua bài toán phân loại tam giác", align="center", size=13,
  italic=True, after=4)
P("`Triangle.classify(int a, int b, int c)`", align="center", size=13, after=60)
P("Học kỳ I — Năm học 2026–2027", align="center", size=12, color=RGBColor(0x59, 0x59, 0x59))
PB()

# ======================================================================
# MUC LUC + MUC TIEU
# ======================================================================

H1("Nội dung chính")
T(["Phần", "Nội dung", "Bạn sẽ làm được gì"],
  [["1", "Đề bài và hiểu chương trình", "Đọc hiểu `classify`, phân biệt hộp trắng / hộp đen"],
   ["2", "Control Flow Graph (CFG)", "Tự vẽ CFG từ mã nguồn theo 5 bước"],
   ["3", "Cyclomatic Complexity (CC)", "Tính CC bằng 3 công thức và hiểu ý nghĩa"],
   ["4", "Đường đi độc lập", "Tìm tập đường cơ sở (basis path)"],
   ["5", "Thiết kế test case", "Chọn dữ liệu cho từng đường, tính Statement / Branch Coverage"],
   ["6", "Kiểm chứng bằng JUnit + JaCoCo", "Chạy test, đọc báo cáo độ bao phủ"],
   ["7", "Tóm tắt", "Bảng tra nhanh + bản trả lời ngắn gọn để nộp bài"],
   ["8", "Câu hỏi tự luyện", "Tự kiểm tra lại kiến thức (có đáp án)"]],
  widths=[1.4, 6.0, 9.3], center_cols=(0,))

BOX("tip", [
    "Hãy đọc theo thứ tự. Mỗi phần dùng kết quả của phần trước: **CFG → CC → đường đi → test case → đo độ bao phủ**.",
    "Sau mỗi phần, thử **che lời giải và tự làm lại** trên giấy — đó là cách nhớ lâu nhất.",
], title="CÁCH DÙNG TÀI LIỆU NÀY")

# ======================================================================
# PHAN 1
# ======================================================================

H1("1. Đề bài và hiểu chương trình")
H2("1.1. Đề bài")
P("Xét chương trình Java sau:")
CODE('''
public class Triangle {

    public static String classify(int a, int b, int c) {

        if (a <= 0 || b <= 0 || c <= 0) {
            return "Invalid";
        }

        if (a + b <= c || a + c <= b || b + c <= a) {
            return "Not a triangle";
        }

        if (a == b && b == c) {
            return "Equilateral";
        } else if (a == b || b == c || a == c) {
            return "Isosceles";
        } else {
            return "Scalene";
        }
    }
}
''', title="Triangle.java")
P("Yêu cầu:")
B("**(1)** Vẽ Control Flow Graph (CFG) cho phương thức `classify`.")
B("**(2)** Tính Cyclomatic Complexity (CC).")
B("**(3)** Liệt kê các đường đi độc lập (independent paths).")
B("**(4)** Thiết kế bộ test case đạt **100% Statement Coverage** và **100% Branch Coverage**.")

H2("1.2. Chương trình làm gì?")
P("Trước khi vẽ bất cứ thứ gì, hãy hiểu chương trình. `classify` nhận độ dài 3 cạnh `a`, `b`, `c` "
  "và trả về loại tam giác. Nó kiểm tra lần lượt qua **3 bước**, bước nào \"bắt\" được thì trả kết quả "
  "ngay (lệnh `return` kết thúc phương thức):")
T(["Bước", "Điều kiện trong code", "Ý nghĩa", "Kết quả"],
  [["1", "`a <= 0 || b <= 0 || c <= 0`", "Có cạnh không dương (≤ 0)", "`\"Invalid\"`"],
   ["2", "`a + b <= c || a + c <= b || b + c <= a`",
    "Vi phạm bất đẳng thức tam giác: tổng 2 cạnh nào đó không lớn hơn cạnh còn lại", "`\"Not a triangle\"`"],
   ["3a", "`a == b && b == c`", "Ba cạnh bằng nhau", "`\"Equilateral\"` (đều)"],
   ["3b", "`a == b || b == c || a == c`", "Có ít nhất hai cạnh bằng nhau", "`\"Isosceles\"` (cân)"],
   ["3c", "(còn lại)", "Ba cạnh khác nhau", "`\"Scalene\"` (thường)"]],
  widths=[1.3, 6.2, 5.6, 3.6], center_cols=(0,))

BOX("note", [
    "**Bất đẳng thức tam giác:** ba đoạn thẳng tạo thành tam giác khi và chỉ khi tổng hai cạnh bất kỳ "
    "**lớn hơn hẳn** cạnh còn lại. Code kiểm tra điều ngược lại (`<=`) để loại bỏ.",
    "Ví dụ `(1, 2, 3)`: `1 + 2 = 3` → ba điểm thẳng hàng, tam giác \"dẹt\" (suy biến) → `Not a triangle`.",
    "**Thứ tự kiểm tra rất quan trọng:** `(5, 1, 1)` có hai cạnh bằng nhau nhưng kết quả là "
    "`Not a triangle`, vì bị chặn ở bước 2 (`1 + 1 <= 5`) trước khi tới bước 3.",
])

H2("1.3. Kiểm thử hộp trắng khác gì hộp đen?")
T(["", "Hộp đen (Lab02)", "Hộp trắng (Lab03)"],
  [["Dựa vào", "Đặc tả / yêu cầu, không nhìn code", "Cấu trúc bên trong của mã nguồn"],
   ["Câu hỏi đặt ra", "Chương trình có làm **đúng yêu cầu** không?",
    "Bộ test đã **chạy qua** những phần nào của code?"],
   ["Kỹ thuật", "Phân vùng tương đương, giá trị biên",
    "CFG, Cyclomatic Complexity, basis path, độ bao phủ"],
   ["Thước đo", "Số vùng / số biên đã kiểm thử", "% câu lệnh, % nhánh được thực thi"]],
  widths=[3.2, 6.7, 6.8])
P("Hai loại bổ sung cho nhau: hộp trắng giúp phát hiện **đoạn code chưa bao giờ được chạy thử**, "
  "còn hộp đen giúp phát hiện **chức năng bị thiếu hoặc hiểu sai yêu cầu**.")

# ======================================================================
# PHAN 2
# ======================================================================

H1("2. Control Flow Graph (CFG)")
H2("2.1. CFG là gì?")
P("**Control Flow Graph** (đồ thị luồng điều khiển) là một đồ thị có hướng mô tả **mọi cách** mà luồng thực thi "
  "có thể đi qua chương trình.")
T(["Thành phần", "Ý nghĩa", "Ký hiệu trong tài liệu"],
  [["Nút (node)", "Một câu lệnh, hoặc một khối lệnh luôn chạy liền nhau (không rẽ nhánh)", "Hình tròn xanh dương"],
   ["Nút quyết định (predicate / decision node)",
    "Nút có **2 cạnh ra** — ứng với điều kiện của `if`, `while`, `for`…", "Hình tròn màu cam"],
   ["Cạnh (edge)", "Luồng điều khiển có thể đi từ nút này sang nút kia", "Mũi tên"],
   ["Nhãn T / F", "Cạnh đi khi điều kiện đúng (True) / sai (False)", "Chữ T xanh lá / F đỏ"],
   ["Nút vào / nút ra", "Nơi bắt đầu và kết thúc phương thức (chỉ có **một** nút ra)", "Nút 1 / nút 10 (vòng tròn kép)"]],
  widths=[4.4, 7.8, 4.5])

H2("2.2. Quy tắc chuyển mã nguồn thành CFG")
T(["Cấu trúc trong code", "Cách vẽ trên CFG"],
  [["Các lệnh tuần tự, không rẽ nhánh", "Gộp thành **một nút**"],
   ["`if (C) { A }`", "Nút quyết định C: cạnh **T → A**, cạnh **F → lệnh ngay sau khối if**"],
   ["`if (C) { A } else { B }`", "Nút quyết định C: cạnh **T → A**, cạnh **F → B**"],
   ["`else if (C2)`", "Là một **nút quyết định mới** nằm trên nhánh F của `if` trước"],
   ["`else` (không có điều kiện)", "**Không** phải nút — nó chỉ là đích của nhánh F"],
   ["`return ...`", "Nút lệnh, có **một cạnh** đi tới nút ra duy nhất"]],
  widths=[6.0, 10.7])

H2("2.3. Vẽ CFG cho classify — từng bước")
H3("Bước 1 — Đánh số các câu lệnh")
P("Gán cho mỗi điều kiện `if` và mỗi lệnh `return` một số. Thêm một nút **exit** cho điểm kết thúc chung.")
CODE('''
public static String classify(int a, int b, int c) {

    if (a <= 0 || b <= 0 || c <= 0) {                   // (1)
        return "Invalid";                               // (2)
    }

    if (a + b <= c || a + c <= b || b + c <= a) {       // (3)
        return "Not a triangle";                        // (4)
    }

    if (a == b && b == c) {                             // (5)
        return "Equilateral";                           // (6)
    } else if (a == b || b == c || a == c) {            // (7)
        return "Isosceles";                             // (8)
    } else {
        return "Scalene";                               // (9)
    }
}                                                       // (10) exit
''', title="classify() sau khi đánh số nút")

H3("Bước 2 — Lập bảng nút")
T(["Nút", "Loại", "Nội dung"],
  [["1", "Quyết định **D1**", "`a <= 0 || b <= 0 || c <= 0`"],
   ["2", "Lệnh", "`return \"Invalid\"`"],
   ["3", "Quyết định **D2**", "`a + b <= c || a + c <= b || b + c <= a`"],
   ["4", "Lệnh", "`return \"Not a triangle\"`"],
   ["5", "Quyết định **D3**", "`a == b && b == c`"],
   ["6", "Lệnh", "`return \"Equilateral\"`"],
   ["7", "Quyết định **D4**", "`a == b || b == c || a == c`"],
   ["8", "Lệnh", "`return \"Isosceles\"`"],
   ["9", "Lệnh", "`return \"Scalene\"`"],
   ["10", "Kết thúc", "Thoát khỏi phương thức"]],
  widths=[1.6, 4.0, 11.1], center_cols=(0,))
P("⇒ Có **N = 10 nút**, trong đó **4 nút quyết định** (1, 3, 5, 7).")

H3("Bước 3 — Nối cạnh")
P("Với mỗi nút quyết định, hỏi hai câu: **\"Nếu đúng thì đi đâu? Nếu sai thì đi đâu?\"** "
  "Với mỗi lệnh `return`: luôn đi tới nút 10.")
T(["Cạnh", "Từ → Đến", "Khi nào", "Cạnh", "Từ → Đến", "Khi nào"],
  [["e1", "1 → 2", "D1 = True", "e8", "7 → 8", "D4 = True"],
   ["e2", "1 → 3", "D1 = False", "e9", "7 → 9", "D4 = False"],
   ["e3", "3 → 4", "D2 = True", "e10", "4 → 10", "return"],
   ["e4", "3 → 5", "D2 = False", "e11", "6 → 10", "return"],
   ["e5", "2 → 10", "return", "e12", "8 → 10", "return"],
   ["e6", "5 → 6", "D3 = True", "e13", "9 → 10", "return"],
   ["e7", "5 → 7", "D3 = False", "", "", ""]],
  widths=[1.6, 2.6, 4.1, 1.6, 2.6, 4.2], center_cols=(0, 1, 3, 4))
P("⇒ Có **E = 13 cạnh**.")

H3("Bước 4 — Vẽ hình")
IMG("cfg.png", 14.5, "Hình 1. Control Flow Graph của Triangle.classify (mức quyết định)")

H3("Bước 5 — Tự kiểm tra hình vừa vẽ")
B("Mỗi nút quyết định có **đúng 2** cạnh ra (một T, một F)? → nút 1, 3, 5, 7: đúng.")
B("Mỗi nút `return` có **đúng 1** cạnh, đi tới nút ra? → nút 2, 4, 6, 8, 9: đúng.")
B("Chỉ có **một** nút ra (nút 10)? → đúng.")
B("Từ nút 1 đi tới được mọi nút, và từ mọi nút đi tới được nút 10? → đúng.")

BOX("ex", [
    "Lần theo hình với đầu vào `(3, 3, 5)`:",
    "• Nút 1: `3<=0`? `3<=0`? `5<=0`? → tất cả sai → **F** → nút 3.",
    "• Nút 3: `6<=5`? `8<=3`? `8<=3`? → tất cả sai → **F** → nút 5.",
    "• Nút 5: `3==3` đúng, nhưng `3==5` sai → cả biểu thức `&&` **F** → nút 7.",
    "• Nút 7: `3==3` đúng → **T** → nút 8 → trả về `\"Isosceles\"` → nút 10.",
    "Đường đi: **1 → 3 → 5 → 7 → 8 → 10**.",
], title="VÍ DỤ — ĐỌC CFG")

H2("2.4. Điều kiện phức và short-circuit")
P("Các điều kiện trong bài có **nhiều vế** nối bằng `||` hoặc `&&`. Java đánh giá chúng theo kiểu "
  "**short-circuit** (đoản mạch):")
B("`X || Y`: nếu `X` đúng thì **dừng luôn**, kết quả là đúng, không cần xét `Y`.")
B("`X && Y`: nếu `X` sai thì **dừng luôn**, kết quả là sai, không cần xét `Y`.")
P("Vì vậy có **hai mức** vẽ CFG:")
IMG("split_d1.png", 15.5, "Hình 2. Nút D1 ở mức quyết định (trái) và khi tách thành điều kiện đơn (phải)")
T(["Mức vẽ", "Mỗi nút quyết định là", "Dùng cho"],
  [["**Mức quyết định** (bài làm dùng)", "Cả biểu thức `if (...)`", "Statement Coverage, Branch Coverage"],
   ["Mức điều kiện đơn", "Từng vế `a <= 0`, `b <= 0`…", "Condition Coverage, MC/DC; công cụ JaCoCo"]],
  widths=[5.4, 5.4, 5.9])
P("Đề bài chỉ hỏi Statement và Branch Coverage nên bài làm vẽ CFG ở **mức quyết định**. "
  "Mức điều kiện đơn được nhắc lại ở phần 3.3 và phần 6.")

BOX("warn", [
    "• Coi `else` là một nút quyết định → đếm thừa nút, CC sai.",
    "• Vẽ mỗi `return` một nút kết thúc riêng → công thức `E − N + 2` cho kết quả sai. Luôn gom về **một** nút ra.",
    "• Quên cạnh F của `if` (khi `if` không có `else`, nhánh F đi tới lệnh ngay sau khối `if`).",
    "• Vẽ nhầm: nhánh F của nút 1 đi thẳng tới nút 5, bỏ qua nút 3.",
])

# ======================================================================
# PHAN 3
# ======================================================================

H1("3. Cyclomatic Complexity (CC)")
H2("3.1. CC là gì và dùng để làm gì?")
P("**Cyclomatic Complexity** — độ phức tạp chu trình, do Thomas McCabe đề xuất năm 1976 — ký hiệu **V(G)**. "
  "Nó là **số đường đi độc lập** trong CFG. Con số này cho biết hai điều:")
B("**Số test case cần có** để phủ tập đường cơ sở (basis path). Đây cũng là **cận trên** "
  "của số test cần để đạt 100% Branch Coverage.")
B("**Độ phức tạp của code**: CC càng lớn, code càng khó hiểu, khó kiểm thử, dễ có lỗi.")
T(["Giá trị CC", "Đánh giá thường dùng"],
  [["1 – 10", "Đơn giản, rủi ro thấp — dễ kiểm thử"],
   ["11 – 20", "Phức tạp vừa, rủi ro trung bình"],
   ["21 – 50", "Phức tạp, rủi ro cao — nên tách hàm"],
   ["> 50", "Rất khó kiểm thử"]],
  widths=[4.0, 12.7], center_cols=(0,))

H2("3.2. Ba cách tính — cả ba phải ra cùng một số")
H3("Cách 1 — Theo số cạnh và số nút")
FORMULA("V(G) = E − N + 2", "E = số cạnh, N = số nút (đồ thị liên thông, 1 nút vào, 1 nút ra)")
P("Với CFG của `classify`: E = 13, N = 10 ⇒ **V(G) = 13 − 10 + 2 = 5**.", align="center")

H3("Cách 2 — Theo số nút quyết định")
FORMULA("V(G) = P + 1", "P = số nút quyết định có 2 nhánh")
P("Các nút quyết định là 1, 3, 5, 7 ⇒ P = 4 ⇒ **V(G) = 4 + 1 = 5**.", align="center")
P("Lưu ý: `else` không có điều kiện nên không được tính. Nếu gặp `switch` có k nhánh thì nút đó đóng góp k − 1.")

H3("Cách 3 — Theo số miền của đồ thị")
FORMULA("V(G) = số miền (region)", "kể cả miền bên ngoài đồ thị")
P("Khi vẽ CFG trên mặt phẳng mà không có cạnh nào cắt nhau, các cạnh chia mặt phẳng thành các miền. "
  "Mỗi miền kín được bao bởi các cạnh; miền bên ngoài cũng tính là một miền.")
IMG("cfg_regions.png", 14.0, "Hình 3. Bốn miền kín R1–R4 và miền ngoài R5 ⇒ V(G) = 5")
P("Nhận xét: mỗi nút quyết định \"mở ra\" thêm đúng một miền kín (R1 ở nút 1, R2 ở nút 3, R3 ở nút 5, R4 ở nút 7). "
  "Vì vậy số miền = P + 1 — ba công thức thực chất là một.")

BOX("tip", [
    "Luôn tính CC bằng **ít nhất 2 cách**. Nếu kết quả khác nhau ⇒ CFG đang vẽ sai (thiếu cạnh, thừa nút, "
    "hoặc có nhiều nút ra).",
])

H2("3.3. Nếu tách điều kiện phức thành điều kiện đơn")
P("Khi vẽ CFG ở mức điều kiện đơn (như Hình 2 bên phải), mỗi vế là một nút quyết định:")
T(["Quyết định", "Các điều kiện đơn", "Số lượng"],
  [["D1", "`a <= 0`,  `b <= 0`,  `c <= 0`", "3"],
   ["D2", "`a + b <= c`,  `a + c <= b`,  `b + c <= a`", "3"],
   ["D3", "`a == b`,  `b == c`", "2"],
   ["D4", "`a == b`,  `b == c`,  `a == c`", "3"],
   ["**Tổng**", "", "**11**"]],
  widths=[2.8, 11.4, 2.5], center_cols=(0, 2))
P("⇒ V(G) = 11 + 1 = **12**. Công cụ JaCoCo cũng báo complexity của `classify` là **12**, vì nó đếm trên bytecode, "
  "nơi mỗi vế là một lệnh rẽ nhánh riêng.")
BOX("note", [
    "Kết quả chính của bài: **CC = 5** (CFG mức quyết định — đúng với yêu cầu Statement/Branch Coverage).",
    "Nếu giảng viên yêu cầu tách điều kiện đơn thì **CC = 12**. Khi không chắc, hãy trình bày cả hai và nói rõ "
    "đang dùng mức nào — như bài làm đã ghi chú.",
    "CC = 5 nằm trong khoảng 1–10 ⇒ `classify` là hàm **đơn giản, dễ kiểm thử**.",
])

# ======================================================================
# PHAN 4
# ======================================================================

H1("4. Đường đi độc lập (Independent / Basis paths)")
H2("4.1. Định nghĩa")
B("**Đường đi (path):** dãy nút từ nút vào (1) tới nút ra (10), đi theo chiều mũi tên.")
B("**Đường đi độc lập:** đường có **ít nhất một cạnh mới** — cạnh chưa xuất hiện trong các đường đã chọn trước đó.")
B("**Tập đường cơ sở (basis set):** tập các đường độc lập có **đúng V(G) đường**. Đi hết tập này thì mọi cạnh, "
  "mọi nút của CFG đều được đi qua.")

H2("4.2. Cách tìm: lần lượt \"lật\" từng nút quyết định")
P("Cách làm dễ nhớ nhất cho chuỗi `if` như bài này:")
B("**Bước 1:** Chọn đường đầu tiên: rẽ **T** ngay tại nút quyết định đầu tiên.")
B("**Bước 2:** Đường tiếp theo: giữ các nút trước ở **F**, rẽ **T** tại nút quyết định kế tiếp.")
B("**Bước 3:** Lặp lại đến nút quyết định cuối. Đường cuối cùng đi **F** ở tất cả các nút.")
B("**Bước 4:** Dừng khi đủ V(G) = 5 đường và kiểm tra đã phủ đủ 13 cạnh.")
T(["Đường", "D1", "D2", "D3", "D4", "Dãy nút", "Cạnh mới", "Kết quả"],
  [["**P1**", "T", "–", "–", "–", "1 → 2 → 10", "e1, e5", "Invalid"],
   ["**P2**", "F", "T", "–", "–", "1 → 3 → 4 → 10", "e2, e3, e10", "Not a triangle"],
   ["**P3**", "F", "F", "T", "–", "1 → 3 → 5 → 6 → 10", "e4, e6, e11", "Equilateral"],
   ["**P4**", "F", "F", "F", "T", "1 → 3 → 5 → 7 → 8 → 10", "e7, e8, e12", "Isosceles"],
   ["**P5**", "F", "F", "F", "F", "1 → 3 → 5 → 7 → 9 → 10", "e9, e13", "Scalene"]],
  widths=[1.5, 0.95, 0.95, 0.95, 0.95, 5.0, 2.7, 3.7], center_cols=(0, 1, 2, 3, 4), size=10)
P("Dấu \"–\" nghĩa là đường đi đã kết thúc trước khi tới nút quyết định đó.", size=10.5, italic=True)
IMG("paths.png", 16.7, "Hình 4. Năm đường đi độc lập P1–P5 (tô đỏ) trên CFG")

H2("4.3. Kiểm tra lại")
B("Mỗi đường đều có cạnh mới (cột \"Cạnh mới\" không rỗng) ⇒ các đường **độc lập**.")
B("Tổng số cạnh mới: 2 + 3 + 3 + 3 + 2 = **13 = E** ⇒ phủ **toàn bộ** cạnh.")
B("Số đường = 5 = **V(G)** ⇒ đây là một **tập đường cơ sở** hoàn chỉnh.")
BOX("note", [
    "Nói chung tập đường cơ sở **không duy nhất** (nhất là khi có vòng lặp), nhưng **số lượng** luôn bằng V(G).",
    "Riêng bài này, CFG không có vòng lặp và các nhánh chỉ gặp lại nhau ở nút 10, nên chương trình có "
    "**đúng 5 đường** từ đầu đến cuối. Tập cơ sở cũng chính là tất cả các đường.",
    "Cả 5 đường đều **khả thi** (tìm được dữ liệu để đi qua). Ở bài khác có thể gặp **đường bất khả thi** "
    "(infeasible path) — đường có trên đồ thị nhưng không dữ liệu nào đi qua được; khi đó ghi rõ và bỏ qua.",
])

# ======================================================================
# PHAN 5
# ======================================================================

H1("5. Thiết kế test case")
H2("5.1. Quy trình 4 bước")
T(["Bước", "Việc cần làm", "Ví dụ với P4"],
  [["1", "Chọn một đường đi", "P4: 1 → 3 → 5 → 7 → 8 → 10"],
   ["2", "Viết **điều kiện đường đi** (path condition): D1..D4 phải đúng/sai thế nào",
    "D1 = F, D2 = F, D3 = F, D4 = T"],
   ["3", "Chọn giá trị đầu vào thỏa điều kiện đó", "(3, 3, 5)"],
   ["4", "Xác định **kết quả mong đợi từ đặc tả** (không phải chạy code rồi chép lại)", "\"Isosceles\""]],
  widths=[1.4, 9.0, 6.3], center_cols=(0,))

H2("5.2. Điều kiện đường đi và dữ liệu được chọn")
T(["Đường", "Điều kiện cần thỏa", "Đầu vào (a, b, c)", "Vì sao chọn"],
  [["P1", "Có ít nhất một cạnh ≤ 0", "(0, 4, 5)", "`a = 0` ⇒ D1 đúng ngay ở vế đầu"],
   ["P2", "Cả 3 cạnh > 0\nVI PHẠM bất đẳng thức tam giác", "(1, 2, 3)", "`1 + 2 = 3` ⇒ `a + b <= c` đúng (tam giác dẹt)"],
   ["P3", "Cạnh > 0, là tam giác\na = b = c", "(3, 3, 3)", "Ba cạnh bằng nhau"],
   ["P4", "Cạnh > 0, là tam giác\nkhông đều, có 2 cạnh bằng nhau", "(3, 3, 5)", "`a = b ≠ c`, và 3 + 3 > 5"],
   ["P5", "Cạnh > 0, là tam giác\nba cạnh khác nhau", "(3, 4, 5)", "Tam giác vuông quen thuộc 3-4-5"]],
  widths=[1.5, 5.8, 3.3, 6.1], center_cols=(0, 2))

H2("5.3. Bảng test case")
T(["ID", "Đường", "Đầu vào (a, b, c)", "Kết quả mong đợi"],
  [["TC01", "P1", "(0, 4, 5)", "Invalid"],
   ["TC02", "P2", "(1, 2, 3)", "Not a triangle"],
   ["TC03", "P3", "(3, 3, 3)", "Equilateral"],
   ["TC04", "P4", "(3, 3, 5)", "Isosceles"],
   ["TC05", "P5", "(3, 4, 5)", "Scalene"]],
  widths=[2.4, 2.4, 5.0, 6.9], center_cols=(0, 1, 2, 3))

H2("5.4. Lần vết (trace) từng test case")
P("Tính giá trị từng vế của mỗi điều kiện để chắc chắn test case đi đúng đường đã định. "
  "Chữ \"bỏ qua\" nghĩa là vế đó không được tính do short-circuit.")
T(["TC", "D1 (nút 1)", "D2 (nút 3)", "D3 (nút 5)", "D4 (nút 7)", "Đường"],
  [["TC01\n(0,4,5)", "0<=0 **T**\n(bỏ qua b, c)\n⇒ **T**", "–", "–", "–", "P1"],
   ["TC02\n(1,2,3)", "F, F, F\n⇒ **F**", "1+2<=3 **T**\n(bỏ qua 2 vế sau)\n⇒ **T**", "–", "–", "P2"],
   ["TC03\n(3,3,3)", "F, F, F\n⇒ **F**", "6<=3 F\n6<=3 F\n6<=3 F\n⇒ **F**", "3==3 T\n3==3 T\n⇒ **T**", "–", "P3"],
   ["TC04\n(3,3,5)", "F, F, F\n⇒ **F**", "6<=5 F\n8<=3 F\n8<=3 F\n⇒ **F**", "3==3 T\n3==5 F\n⇒ **F**",
    "3==3 **T**\n(bỏ qua 2 vế sau)\n⇒ **T**", "P4"],
   ["TC05\n(3,4,5)", "F, F, F\n⇒ **F**", "7<=5 F\n8<=4 F\n9<=3 F\n⇒ **F**", "3==4 F\n(bỏ qua vế sau)\n⇒ **F**",
    "3==4 F\n4==5 F\n3==5 F\n⇒ **F**", "P5"]],
  widths=[2.1, 2.9, 3.4, 3.0, 3.4, 1.9], center_cols=(0, 1, 2, 3, 4, 5), size=10)

H2("5.5. 100% Statement Coverage")
P("**Statement Coverage** (độ phủ câu lệnh, còn gọi là phủ cấp 1): mỗi câu lệnh được thực thi **ít nhất một lần**.")
FORMULA("Statement Coverage = số câu lệnh đã chạy / tổng số câu lệnh × 100%", width=16.2, size=12.5)
T(["Nút / câu lệnh", "TC01", "TC02", "TC03", "TC04", "TC05"],
  [["(1) if D1", "✔", "✔", "✔", "✔", "✔"],
   ["(2) return \"Invalid\"", "✔", "", "", "", ""],
   ["(3) if D2", "", "✔", "✔", "✔", "✔"],
   ["(4) return \"Not a triangle\"", "", "✔", "", "", ""],
   ["(5) if D3", "", "", "✔", "✔", "✔"],
   ["(6) return \"Equilateral\"", "", "", "✔", "", ""],
   ["(7) else if D4", "", "", "", "✔", "✔"],
   ["(8) return \"Isosceles\"", "", "", "", "✔", ""],
   ["(9) return \"Scalene\"", "", "", "", "", "✔"]],
  widths=[6.2, 2.1, 2.1, 2.1, 2.1, 2.1], center_cols=(1, 2, 3, 4, 5))
P("⇒ 9 / 9 câu lệnh được thực thi ⇒ **Statement Coverage = 100%**.", bold=False)
BOX("note", [
    "**Vì sao cần tối thiểu 5 test case?** Mỗi lần gọi `classify` chỉ dừng ở **đúng một** lệnh `return`. "
    "Có 5 lệnh `return` khác nhau (nút 2, 4, 6, 8, 9) ⇒ cần ít nhất 5 lần chạy mới chạm hết.",
])

H2("5.6. 100% Branch Coverage")
P("**Branch Coverage** (độ phủ nhánh, còn gọi là Decision Coverage hay phủ cấp 2): mỗi nút quyết định phải "
  "nhận **cả hai kết quả** True và False, ít nhất một lần mỗi kết quả. Nói cách khác: **mọi cạnh ra** của nút "
  "quyết định đều được đi qua.")
FORMULA("Branch Coverage = số nhánh đã đi / tổng số nhánh × 100%", width=16.2, size=12.5)
P("Có 4 nút quyết định × 2 nhánh = **8 nhánh**:")
T(["Nhánh", "TC01", "TC02", "TC03", "TC04", "TC05"],
  [["D1 = True  (1 → 2)", "✔", "", "", "", ""],
   ["D1 = False (1 → 3)", "", "✔", "✔", "✔", "✔"],
   ["D2 = True  (3 → 4)", "", "✔", "", "", ""],
   ["D2 = False (3 → 5)", "", "", "✔", "✔", "✔"],
   ["D3 = True  (5 → 6)", "", "", "✔", "", ""],
   ["D3 = False (5 → 7)", "", "", "", "✔", "✔"],
   ["D4 = True  (7 → 8)", "", "", "", "✔", ""],
   ["D4 = False (7 → 9)", "", "", "", "", "✔"]],
  widths=[6.2, 2.1, 2.1, 2.1, 2.1, 2.1], center_cols=(1, 2, 3, 4, 5))
P("⇒ 8 / 8 nhánh được đi qua ⇒ **Branch Coverage = 100%**, dùng **cùng** 5 test case.")

BOX("ex", [
    "Nếu chỉ dùng **TC03 và TC05**:",
    "• Câu lệnh đã chạy: TC03 → 1, 3, 5, 6; TC05 → 1, 3, 5, 7, 9 ⇒ {1, 3, 5, 6, 7, 9} = 6 / 9 ⇒ "
    "**Statement Coverage ≈ 66,7%**.",
    "• Nhánh đã đi: D1=F, D2=F, D3=T, D3=F, D4=F = 5 / 8 ⇒ **Branch Coverage = 62,5%**.",
], title="VÍ DỤ — ĐỘ PHỦ CHƯA ĐỦ")

H2("5.7. Quan hệ giữa Statement và Branch Coverage")
P("**100% Branch Coverage ⇒ 100% Statement Coverage**, nhưng **chiều ngược lại không đúng**. Ví dụ kinh điển:")
CODE('''
int abs(int x) {
    if (x < 0)
        x = -x;
    return x;
}
''', title="Ví dụ: if không có else")
B("Chỉ một test `x = -3`: chạy qua **mọi** câu lệnh ⇒ Statement Coverage = 100%.")
B("Nhưng nhánh `x < 0` = False (trường hợp `x >= 0`) **chưa bao giờ** được thử ⇒ Branch Coverage = 50%.")
P("Trong `classify`, mỗi nhánh đều dẫn tới một câu lệnh riêng, nên hai tiêu chí tình cờ cần **cùng** 5 test case.")

T(["Tiêu chí", "Yêu cầu", "Với classify cần"],
  [["Statement Coverage", "Mỗi câu lệnh chạy ≥ 1 lần", "5 test (TC01–TC05)"],
   ["Branch / Decision Coverage", "Mỗi quyết định nhận cả T và F", "5 test (TC01–TC05)"],
   ["Basis Path Coverage", "Đi qua V(G) đường độc lập", "5 test (TC01–TC05)"],
   ["Condition Coverage (mức đơn)", "Mỗi vế điều kiện nhận cả T và F", "Cần thêm test (xem mục 6.4)"]],
  widths=[5.2, 6.3, 5.2], caption="Bảng so sánh các tiêu chí — tiêu chí bên dưới mạnh hơn tiêu chí bên trên")

BOX("warn", [
    "• **Chép kết quả chạy chương trình làm \"kết quả mong đợi\".** Nếu code sai, test sẽ vẫn \"xanh\" và "
    "không phát hiện được lỗi. Kết quả mong đợi phải suy ra từ **đặc tả** (định nghĩa tam giác).",
    "• Chọn dữ liệu \"nghĩ là\" đi đường P4 nhưng thực ra đi P2, ví dụ `(5, 1, 1)`. Luôn **lần vết** như mục 5.4.",
    "• Nghĩ rằng đạt 100% Statement Coverage là đủ — Branch Coverage mới là yêu cầu tối thiểu thường dùng.",
])

# ======================================================================
# PHAN 6
# ======================================================================

H1("6. Kiểm chứng bằng JUnit và JaCoCo")
H2("6.1. Cấu trúc project Lab03")
CODE('''
Lab03
├── pom.xml                                  <- cau hinh Maven, JUnit 5, JaCoCo
├── README.md
├── docs
│   └── Bai01-CFG-va-do-bao-phu-Triangle.md  <- bai lam dang Markdown
└── src
    ├── main/java/Triangle.java              <- ma nguon can kiem thu
    └── test/java/TriangleTest.java          <- 5 test case TC01 - TC05
''')

H2("6.2. Đọc hiểu mã test JUnit")
CODE('''
import static org.junit.jupiter.api.Assertions.assertEquals;
import org.junit.jupiter.api.Test;

public class TriangleTest {

    @Test
    void TC01_P1_canhKhongDuong_invalid() {
        assertEquals("Invalid", Triangle.classify(0, 4, 5));
    }

    @Test
    void TC04_P4_haiCanhBangNhau_isosceles() {
        assertEquals("Isosceles", Triangle.classify(3, 3, 5));
    }
    // ... TC02, TC03, TC05 tuong tu
}
''', title="TriangleTest.java (trích)")
B("`@Test`: đánh dấu một phương thức là một test case để JUnit tự chạy.")
B("`assertEquals(expected, actual)`: so sánh **kết quả mong đợi** (tham số đầu) với **kết quả thực tế** "
  "(tham số sau). Khác nhau ⇒ test **FAIL**.")
B("Tên test gồm: mã test case – đường đi – ý nghĩa – kết quả. Đọc tên là biết test kiểm tra gì.")

H2("6.3. JaCoCo — công cụ đo độ bao phủ")
P("**JaCoCo** (Java Code Coverage) theo dõi xem khi chạy test, dòng code / nhánh nào đã được thực thi. "
  "Nó được khai báo trong `pom.xml`:")
CODE('''
<plugin>
    <groupId>org.jacoco</groupId>
    <artifactId>jacoco-maven-plugin</artifactId>
    <version>0.8.12</version>
    <executions>
        <!-- gan "may do" vao JVM truoc khi chay test -->
        <execution>
            <goals><goal>prepare-agent</goal></goals>
        </execution>
        <!-- xuat bao cao HTML ngay sau pha test -->
        <execution>
            <id>report</id>
            <phase>test</phase>
            <goals><goal>report</goal></goals>
        </execution>
    </executions>
</plugin>
''', title="pom.xml (trích)")
P("Cách chạy (mở terminal trong thư mục `Lab03`):")
CODE('''
mvn test
# Ket qua test : target/surefire-reports/
# Bao cao do phu: target/site/jacoco/index.html  (mo bang trinh duyet)
''')
P("Trong báo cáo, bấm **default → Triangle → classify(int, int, int)** để xem từng dòng code được tô màu:")
T(["Màu", "Ý nghĩa"],
  [["Xanh lá", "Dòng đã được thực thi, mọi nhánh trên dòng đó đều đã đi"],
   ["Vàng", "Dòng đã chạy nhưng **chỉ đi một phần** các nhánh (JaCoCo hiện hình thoi vàng ở lề)"],
   ["Đỏ", "Dòng chưa bao giờ được thực thi"]],
  widths=[3.0, 13.7])

H2("6.4. Kết quả đo và cách hiểu đúng")
T(["Chỉ số (phương thức classify)", "Kết quả", "Ý nghĩa"],
  [["Tests", "5 run, 0 failure", "Cả 5 test cho đúng kết quả mong đợi"],
   ["Lines", "9 / 9  (100%)", "**Statement Coverage = 100%** ✔"],
   ["Instructions", "46 / 46  (100%)", "Mọi lệnh bytecode của `classify` đều đã chạy"],
   ["Branches", "16 / 22  (72%)", "Đo ở **mức điều kiện đơn**: 11 vế × 2 = 22 nhánh"],
   ["Complexity", "12", "Trùng với CC của CFG mức điều kiện đơn (mục 3.3)"]],
  widths=[5.0, 3.8, 7.9], center_cols=(1,))
P("Ở mức cả lớp `Triangle`, báo cáo có thể thấp hơn vì hàm `main` (chỉ để minh họa) không được test gọi tới. "
  "Hãy nhìn vào dòng của phương thức `classify`.", size=11, italic=True)

BOX("warn", [
    "**\"Branches 16/22\" KHÔNG có nghĩa là bài làm chưa đạt 100% Branch Coverage.**",
    "Branch Coverage theo lý thuyết (mục 5.6) tính trên **4 quyết định = 8 nhánh** ⇒ đã đạt 8/8. "
    "JaCoCo thì coi **mỗi vế** của `||` / `&&` là một nhánh riêng ⇒ 22 nhánh. Đo như vậy gần với "
    "**Condition Coverage**, một tiêu chí mạnh hơn yêu cầu của đề.",
], title="ĐỪNG NHẦM")

P("6 nhánh JaCoCo còn thiếu là những vế **chưa bao giờ nhận giá trị true** (do short-circuit, hoặc do dữ liệu "
  "chưa chạm tới). Nếu muốn JaCoCo báo 22/22, có thể thêm các test sau (đã chạy kiểm chứng):")
T(["Vế chưa nhận true", "Test bổ sung (a, b, c)", "Kết quả"],
  [["`b <= 0` (D1)", "(4, 0, 5)", "Invalid"],
   ["`c <= 0` (D1)", "(4, 5, 0)", "Invalid"],
   ["`a + c <= b` (D2)", "(1, 3, 2)", "Not a triangle"],
   ["`b + c <= a` (D2)", "(3, 1, 2)", "Not a triangle"],
   ["`b == c` (D4)", "(5, 3, 3)", "Isosceles"],
   ["`a == c` (D4)", "(3, 5, 3)", "Isosceles"]],
  widths=[5.6, 5.6, 5.5], center_cols=(1, 2), caption="Phần mở rộng — đề bài không bắt buộc")

# ======================================================================
# PHAN 7
# ======================================================================

H1("7. Tóm tắt")
H2("7.1. Bảng tra nhanh")
T(["Khái niệm", "Cách làm / Công thức", "Kết quả Lab03"],
  [["CFG", "Đánh số lệnh → bảng nút → nối cạnh T/F → gom `return` về 1 nút ra",
    "N = 10 nút, E = 13 cạnh,\n4 nút quyết định"],
   ["Cyclomatic Complexity", "V(G) = E − N + 2 = P + 1 = số miền", "**CC = 5**\n(mức điều kiện đơn: 12)"],
   ["Đường đi độc lập", "Lật từng nút quyết định; mỗi đường thêm ≥ 1 cạnh mới", "5 đường P1 – P5"],
   ["Test case", "Mỗi đường một test; kết quả mong đợi lấy từ đặc tả", "TC01 – TC05"],
   ["Statement Coverage", "Số lệnh đã chạy / tổng số lệnh", "9/9 = **100%**"],
   ["Branch Coverage", "Số nhánh đã đi / tổng số nhánh", "8/8 = **100%**"],
   ["Công cụ", "JUnit 5 (chạy test) + JaCoCo (đo độ phủ)", "5/5 pass, Lines 100%"]],
  widths=[3.9, 7.8, 5.0])

H2("7.2. Bản trả lời ngắn gọn (để nộp bài)")
BOX("note", [
    "**Câu 1 — CFG:** 10 nút (1, 3, 5, 7 là nút quyết định; 2, 4, 6, 8, 9 là các lệnh return; 10 là nút ra), "
    "13 cạnh — vẽ như Hình 1.",
    "**Câu 2 — CC:** V(G) = E − N + 2 = 13 − 10 + 2 = 5; kiểm tra: P + 1 = 4 + 1 = 5; số miền = 5. ⇒ **CC = 5**.",
    "**Câu 3 — Đường đi độc lập:** P1: 1-2-10;  P2: 1-3-4-10;  P3: 1-3-5-6-10;  P4: 1-3-5-7-8-10;  "
    "P5: 1-3-5-7-9-10.",
    "**Câu 4 — Test case:** TC01 (0,4,5) → Invalid;  TC02 (1,2,3) → Not a triangle;  TC03 (3,3,3) → Equilateral;  "
    "TC04 (3,3,5) → Isosceles;  TC05 (3,4,5) → Scalene.  Bộ 5 test này phủ 9/9 câu lệnh và 8/8 nhánh ⇒ "
    "**100% Statement Coverage và 100% Branch Coverage**.",
], title="ĐÁP ÁN RÚT GỌN")

# ======================================================================
# PHAN 8
# ======================================================================

PB()
H1("8. Câu hỏi tự luyện")
P("Tự làm trước, sau đó mới xem đáp án ở trang sau.", italic=True)
QUESTIONS = [
    "Đầu vào `(2, 2, 3)` đi theo đường nào? Kết quả là gì?",
    "Đầu vào `(5, 1, 1)` có hai cạnh bằng nhau. Kết quả có phải `Isosceles` không? Vì sao?",
    "Nếu chỉ dùng TC03 và TC05 thì Statement Coverage và Branch Coverage là bao nhiêu?",
    "Có thể đạt 100% Statement Coverage cho `classify` với ít hơn 5 test case không?",
    "Thêm lệnh `if (a > 100) return \"Too big\";` ngay sau khối kiểm tra `Invalid`. N, E và CC mới là bao nhiêu?",
    "Tìm một đầu vào làm D2 đúng nhờ **vế thứ ba** `b + c <= a` (hai vế đầu sai).",
    "Vì sao `else` cuối cùng (trả về `Scalene`) không được tính là nút quyết định?",
    "JaCoCo báo Branches = 16/22. Bộ test có đạt yêu cầu 100% Branch Coverage của đề không?",
    "Vì sao không nên lấy kết quả chạy chương trình để điền vào cột \"kết quả mong đợi\"?",
]
for i, q in enumerate(QUESTIONS, 1):
    P(f"**Câu {i}.** {q}", after=4)

PB()
H2("Đáp án")
ANSWERS = [
    "D1: F (các cạnh dương). D2: `4<=3` F, `5<=2` F, `5<=2` F ⇒ F. D3: `2==2` T nhưng `2==3` F ⇒ F. "
    "D4: `2==2` T ⇒ T. Đường **P4 (1-3-5-7-8-10)**, kết quả **Isosceles**.",
    "**Không.** D2: `6<=1` F, `6<=1` F, `1+1=2<=5` **T** ⇒ trả về **Not a triangle** (đường P2). "
    "Bước kiểm tra tam giác đứng trước bước phân loại cân/đều.",
    "Statement: {1, 3, 5, 6, 7, 9} = 6/9 ≈ **66,7%**. Branch: D1=F, D2=F, D3=T, D3=F, D4=F = 5/8 = **62,5%**.",
    "**Không.** Mỗi lần chạy chỉ đi tới đúng một trong 5 lệnh `return`, nên cần ít nhất 5 lần chạy.",
    "Thêm 1 nút quyết định X và 1 nút return Y: **N = 12**. Cạnh 1→3 bị thay bằng 1→X, X→Y, X→3, Y→10: "
    "**E = 13 − 1 + 4 = 16**. **CC = 16 − 12 + 2 = 6** (kiểm tra: P + 1 = 5 + 1 = 6).",
    "Ví dụ **(3, 1, 2)**: `3+1<=2` F, `3+2<=1` F, `1+2<=3` **T** ⇒ Not a triangle. Hoặc (10, 2, 3).",
    "`else` không có điều kiện để kiểm tra — nó chỉ là **đích của nhánh F** của nút 7. Nút quyết định phải có "
    "điều kiện và 2 cạnh ra.",
    "**Có.** Branch Coverage của đề tính trên 4 quyết định (8 nhánh) và đã đạt 8/8. JaCoCo đếm theo từng vế "
    "điều kiện (22 nhánh), gần với Condition Coverage — một tiêu chí mạnh hơn.",
    "Nếu code có lỗi, kết quả chạy cũng sai theo ⇒ test vẫn pass và lỗi bị che mất. Kết quả mong đợi phải lấy từ "
    "**đặc tả** (test oracle), độc lập với code.",
]
for i, a in enumerate(ANSWERS, 1):
    P(f"**Câu {i}.** {a}", after=5)

BOX("tip", [
    "Muốn luyện thêm? Hãy tự làm lại toàn bộ Lab03 cho hàm tính tiền điện ở **Lab01** (`ElectricityBill`): "
    "vẽ CFG (có vòng lặp `for`), tính CC, tìm đường đi độc lập và thiết kế test case.",
], title="LUYỆN THÊM")

doc.core_properties.title = "Lab03 — Kiểm thử hộp trắng: Tài liệu ôn tập"
doc.core_properties.subject = "CFG, Cyclomatic Complexity, Basis path, Statement & Branch Coverage"
doc.core_properties.author = "hoangnguyen3101"

doc.save(OUT)
print("Da tao:", os.path.normpath(OUT))
